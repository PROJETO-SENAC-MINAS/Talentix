"""Extração e interpretação local de currículos PDF/DOCX.

A análise é deliberadamente determinística: extrai texto do documento e converte
campos comuns em uma prévia estruturada. O usuário sempre revisa e confirma os
dados antes de qualquer alteração no perfil.
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Iterable

MAX_TEXT_CHARS = 100_000
MAX_PDF_PAGES = 30
PARSER_VERSION = 3

_SECOES = {
    "resumo": {"resumo", "perfil", "perfil profissional", "objetivo", "objetivo profissional", "sobre mim"},
    "experiencias": {"experiencia", "experiencias", "experiencia profissional", "experiencias profissionais", "historico profissional"},
    "formacoes": {"formacao", "formacoes", "formacao academica", "formacoes academicas", "formacao academica e cursos", "educacao", "escolaridade"},
    "certificados": {"certificados", "certificacoes", "cursos", "cursos e certificacoes", "cursos e certificados", "cursos complementares", "qualificacoes"},
    "projetos": {"projeto", "projetos", "projetos pessoais", "projetos academicos", "projetos profissionais", "principais projetos", "projetos relevantes", "portfolio", "portfolio de projetos"},
    "habilidades": {"habilidades", "competencias", "competencias tecnicas", "skills", "tecnologias"},
    "idiomas": {"idiomas", "linguas", "languages"},
}

_CARGO_RE = re.compile(
    r"\b(desenvolvedor|developer|analista|assistente|auxiliar|engenheir[oa]|estagiari[oa]|"
    r"aprendiz|programador|designer|coordenador|gerente|tecnic[oa]|consultor|especialista|"
    r"administrador|arquiteto|cientista|suporte)\b",
    re.IGNORECASE,
)
_INSTITUICAO_RE = re.compile(
    r"\b(universidade|faculdade|instituto|escola|colegio|senac|senai|ifmg|ifsp|ifrs|"
    r"fundacao|coursera|alura|udemy|dio|fiap|puc|ufmg|ufrj|usp|estacio|anhanguera)\b",
    re.IGNORECASE,
)
_NIVEL_FORMACAO = (
    ("doutorado", "Doutorado"),
    ("mestrado", "Mestrado"),
    ("pos-graduacao", "Pós-graduação"),
    ("pos graduacao", "Pós-graduação"),
    ("graduacao", "Graduação"),
    ("bacharel", "Graduação"),
    ("licenciatura", "Graduação"),
    ("tecnologo", "Graduação"),
    ("tecnico", "Técnico"),
    ("ensino medio", "Ensino Médio"),
    ("ensino fundamental / medio", "Ensino Médio"),
    ("ensino fundamental", "Ensino Fundamental"),
    ("curso livre", "Curso livre"),
)
_NIVEIS_IDIOMA = {
    "nativo": "Nativo",
    "fluente": "Fluente",
    "avancado": "Avançado",
    "intermediario": "Intermediário",
    "basico": "Básico",
}
_IDIOMAS_COMUNS = (
    "Português", "Inglês", "Espanhol", "Francês", "Alemão", "Italiano",
    "Mandarim", "Japonês", "Coreano", "Libras",
)


def _normalizar(valor: str) -> str:
    texto = unicodedata.normalize("NFKD", valor or "")
    return "".join(c for c in texto if not unicodedata.combining(c)).lower().strip()


def _limpar_linha(valor: str) -> str:
    valor = re.sub(r"^[\s•●▪◦·‣►✓✔–—-]+", "", valor or "")
    return re.sub(r"\s+", " ", valor).strip()


def extrair_texto_curriculo(caminho: str | Path) -> str:
    """Extrai texto de PDF ou DOCX sem executar conteúdo do documento."""
    caminho = Path(caminho)
    extensao = caminho.suffix.lower()

    if extensao == ".pdf":
        from pypdf import PdfReader

        leitor = PdfReader(str(caminho))
        if leitor.is_encrypted:
            try:
                leitor.decrypt("")
            except Exception as exc:
                raise ValueError("PDF protegido por senha não pode ser analisado.") from exc
        paginas = []
        for pagina in leitor.pages[:MAX_PDF_PAGES]:
            paginas.append(pagina.extract_text() or "")
        texto = "\n".join(paginas)
    elif extensao == ".docx":
        from docx import Document

        documento = Document(str(caminho))
        partes = [p.text for p in documento.paragraphs if p.text and p.text.strip()]
        for tabela in documento.tables:
            for linha in tabela.rows:
                celulas = [c.text.strip() for c in linha.cells if c.text and c.text.strip()]
                if celulas:
                    partes.append(" | ".join(celulas))
        texto = "\n".join(partes)
    elif extensao == ".doc":
        raise ValueError("Arquivos .doc podem ser armazenados, mas a importação automática exige PDF ou DOCX.")
    else:
        raise ValueError("Formato de currículo não suportado para análise automática.")

    texto = texto.replace("\x00", " ")
    texto = "\n".join(re.sub(r"[ \t]+", " ", linha).strip() for linha in texto.splitlines())
    texto = re.sub(r"\n{3,}", "\n\n", texto).strip()
    if len(texto) < 30:
        raise ValueError("Não foi possível extrair texto suficiente do currículo.")
    return texto[:MAX_TEXT_CHARS]


def _identificar_secao(linha: str) -> str | None:
    chave = _normalizar(linha).strip(' :–—-|•·')
    chave = re.sub(r'^\d+[.)]\s*', '', chave)
    exata = next((nome for nome, aliases in _SECOES.items() if chave in aliases), None)
    if exata:
        return exata
    # Cabeçalhos compostos do currículo não são títulos de cursos. Aceitamos
    # somente combinações completas de rótulos, não descrições com essas palavras.
    chave = re.sub(r'\s*[/|•·–—-]\s*', ' ', chave)
    chave = re.sub(r'\s+', ' ', chave).strip()
    if re.fullmatch(r'(?:experiencia pratica\s+(?:e\s+)?)?(?:principais\s+)?projetos(?:\s+(?:pessoais|academicos|profissionais|praticos|relevantes))*', chave):
        return 'projetos'
    if chave == 'experiencia pratica':
        return 'projetos'
    if re.fullmatch(r'(?:principais\s+)?(?:cursos|certificados|certificacoes|capacitacoes)(?:\s+(?:e\s+)?(?:cursos|certificados|certificacoes|complementares|profissionais|adicionais|relevantes))*', chave):
        return 'certificados'
    if re.fullmatch(r'(?:formacao|formacoes)(?:\s+(?:academica|academicas|educacional|educacionais|e|escolaridade))*', chave):
        return 'formacoes'
    return None


def _separar_secoes(linhas: list[str]) -> dict[str, list[str]]:
    secoes: dict[str, list[str]] = {"cabecalho": []}
    atual = "cabecalho"
    for linha in linhas:
        encontrada = _identificar_secao(linha)
        if encontrada:
            atual = encontrada
            secoes.setdefault(atual, [])
            continue
        secoes.setdefault(atual, []).append(linha)
    return secoes


def _url(texto: str, dominio: str) -> str | None:
    padrao = rf"(https?://)?(?:www\.)?{re.escape(dominio)}[^\s,;)>]+"
    achado = re.search(padrao, texto, re.IGNORECASE)
    if not achado:
        return None
    valor = achado.group(0).rstrip(".,;")
    return valor if valor.startswith(("http://", "https://")) else f"https://{valor}"


def _portfolio(texto: str) -> str | None:
    urls = re.findall(r"https?://[^\s,;)>]+|www\.[^\s,;)>]+", texto, re.IGNORECASE)
    for valor in urls:
        low = valor.lower()
        if "linkedin.com" not in low and "github.com" not in low:
            return valor.rstrip(".,;")
    return None


def _localizacao(linhas: Iterable[str]) -> tuple[str | None, str | None]:
    for linha in linhas:
        achado = re.search(r"\b([A-Za-zÀ-ÿ][A-Za-zÀ-ÿ .'-]{2,60}?)\s*[-,/|]\s*([A-Z]{2})\b", linha)
        if achado:
            cidade = achado.group(1).strip(" ,-")
            if len(cidade.split()) <= 7:
                return cidade, achado.group(2).upper()
    return None, None


def _periodo(texto: str) -> tuple[str | None, str | None, bool]:
    normal = _normalizar(texto)
    anos = re.findall(r"\b((?:19|20)\d{2})\b", normal)
    atual = bool(re.search(r"\b(atual|presente|current|hoje)\b", normal))
    # Currículos também escrevem "conclusão 2027, início 2025". A ordem em
    # que os anos aparecem no texto não pode gerar um período impossível.
    anos_ordenados = sorted({int(ano) for ano in anos})
    inicio = f"{anos_ordenados[0]}-01-01" if anos_ordenados else None
    fim = None if atual else (f"{anos_ordenados[-1]}-01-01" if len(anos_ordenados) > 1 else None)
    return inicio, fim, atual


def _inferir_nivel_formacao(texto: str, certificado: bool = False) -> str | None:
    if certificado:
        return "Certificado"
    normal = _normalizar(texto)
    for termo, nivel in _NIVEL_FORMACAO:
        if termo in normal:
            return nivel
    return None


def _parse_formacoes(linhas: list[str], certificado: bool = False) -> list[dict]:
    resultado: list[dict] = []
    vistos: set[tuple[str, str]] = set()

    for indice, linha in enumerate(linhas):
        partes = [_limpar_linha(p) for p in re.split(r"\s*[|•·]\s*|\s+[-–—]\s+", linha) if _limpar_linha(p)]
        # Datas/status isolados podem complementar um registro, mas outra
        # formação nunca deve fornecer seu período ou nível ao item vizinho.
        def metadado(valor):
            return bool(re.fullmatch(r"[\d\s/.,:()–—|\-]+|\d+\s*(?:h|horas)|(?:conclusao|inicio|previsao|cursando|concluido|em andamento|carga horaria|horas|h)\b.*", _normalizar(valor)))
        vizinhas = linhas[max(0, indice - 1):indice] + linhas[indice + 1:indice + 2]
        contexto = " | ".join([linha] + [v for v in vizinhas if metadado(v)])
        nivel = _inferir_nivel_formacao(linha, certificado)
        instituicao = next((p for p in partes if _INSTITUICAO_RE.search(_normalizar(p))), None)
        curso = next((p for p in partes if p != instituicao and not metadado(p)), None)
        if instituicao and not curso:
            curso = next((v for v in vizinhas if not metadado(v) and not _INSTITUICAO_RE.search(_normalizar(v))), None)
            nivel = _inferir_nivel_formacao(curso or linha, certificado)
        if not instituicao and nivel:
            conteudo = [p for p in partes if not metadado(p)]
            if len(conteudo) == 2:
                curso, instituicao = conteudo

        if not instituicao or not curso:
            continue
        if len(instituicao) > 200 or len(curso) > 200:
            continue

        chave = (_normalizar(instituicao), _normalizar(curso))
        if chave in vistos:
            continue
        vistos.add(chave)

        # Prefere o período da própria linha. Usa as linhas vizinhas somente
        # quando instituição/curso e datas foram separados pelo PDF/DOCX.
        # Um ano no título (Office 2016, por exemplo) não é uma data de estudo.
        fonte_periodo = ' | '.join(p for p in partes if metadado(p))
        if not fonte_periodo:
            fonte_periodo = ' | '.join(v for v in vizinhas if metadado(v))
        inicio, fim, _ = _periodo(fonte_periodo)
        resultado.append({
            "instituicao": instituicao,
            "curso": curso,
            "nivel": nivel,
            "data_inicio": inicio,
            "data_conclusao": fim,
            "status": "Cursando" if re.search(r"\b(cursando|em andamento)\b", _normalizar(contexto)) else None,
        })
    return resultado[:20]


def _parse_projetos(linhas: list[str]) -> list[dict]:
    """Reconhece títulos explícitos, preservando as linhas de descrição juntas."""
    resultado = []
    for linha in linhas:
        url = re.search(r"https?://[^\s|<>]+", linha)
        sem_url = re.sub(r"https?://[^\s|<>]+", "", linha).strip(" |–—-")
        partes = re.split(r"\s*[|•·]\s*|\s+[-–—]\s+|:\s+", sem_url, maxsplit=1)
        titulo = partes[0].strip()
        novo = bool(titulo and (not resultado or len(partes) > 1 or (len(titulo) <= 80 and not titulo.endswith('.'))))
        if novo:
            resultado.append({"titulo": titulo[:200], "descricao": partes[1].strip()[:5000] if len(partes) > 1 else None,
                              "url": url.group(0).rstrip('.,;')[:300] if url else None})
        elif resultado:
            if sem_url:
                resultado[-1]["descricao"] = "\n".join(filter(None, [resultado[-1]["descricao"], sem_url]))[:5000]
            if url and not resultado[-1]["url"]:
                resultado[-1]["url"] = url.group(0).rstrip('.,;')[:300]
    return resultado[:20]


def _parse_experiencias(linhas: list[str]) -> list[dict]:
    resultado: list[dict] = []
    vistos: set[tuple[str, str, str]] = set()

    for indice, linha in enumerate(linhas):
        inicio, fim, atual = _periodo(linha)
        if not inicio:
            continue

        anteriores = [_limpar_linha(x) for x in linhas[max(0, indice - 2):indice] if _limpar_linha(x)]
        partes = [_limpar_linha(p) for p in re.split(r"\s*[|•·]\s*|\s+[-–—]\s+", linha) if _limpar_linha(p)]
        partes_sem_data = [p for p in partes if not re.search(r"\b(?:19|20)\d{2}\b|\b(atual|presente)\b", _normalizar(p))]
        candidatos = anteriores + partes_sem_data

        cargo = next((x for x in candidatos if _CARGO_RE.search(_normalizar(x))), None)
        empresa = next((x for x in candidatos if x != cargo and 1 < len(x) <= 150), None)

        if not cargo and candidatos:
            cargo = candidatos[-1]
        if not empresa and len(candidatos) >= 2:
            empresa = candidatos[-2]

        if not cargo or not empresa or cargo == empresa:
            continue

        chave = (_normalizar(empresa), _normalizar(cargo), inicio)
        if chave in vistos:
            continue
        vistos.add(chave)

        descricao = None
        if indice + 1 < len(linhas) and not _periodo(linhas[indice + 1])[0]:
            proxima = _limpar_linha(linhas[indice + 1])
            if proxima and not _CARGO_RE.fullmatch(_normalizar(proxima)):
                descricao = proxima[:2000]

        resultado.append({
            "empresa": empresa[:150],
            "cargo": cargo[:150],
            "descricao": descricao,
            "data_inicio": inicio,
            "data_fim": fim,
            "atual": atual,
        })
    return resultado[:20]


def _parse_habilidades(texto: str, linhas_secao: list[str], catalogo: Iterable[str]) -> list[dict]:
    nomes: list[str] = []
    vistos: set[str] = set()

    for nome in catalogo:
        nome = (nome or "").strip()
        normal = _normalizar(nome)
        if not nome or (len(normal) < 2 and normal not in {"r"}):
            continue
        if re.search(rf"(?<!\w){re.escape(normal)}(?!\w)", _normalizar(texto)):
            if normal not in vistos:
                vistos.add(normal)
                nomes.append(nome)

    for linha in linhas_secao:
        for parte in re.split(r"\s*[|,;•·]\s*", linha):
            nome = _limpar_linha(parte)
            normal = _normalizar(nome)
            if (
                1 < len(nome) <= 80
                and not re.search(r"\b(?:19|20)\d{2}\b", nome)
                and "http" not in normal
                and normal not in vistos
            ):
                vistos.add(normal)
                nomes.append(nome)

    return [{"nome": nome, "nivel": None} for nome in nomes[:50]]


def _parse_idiomas(texto: str, linhas_secao: list[str], catalogo: Iterable[str]) -> list[dict]:
    conhecidos = [x for x in catalogo if x] or list(_IDIOMAS_COMUNS)
    base = "\n".join(linhas_secao) if linhas_secao else texto
    normal_base = _normalizar(base)
    resultado = []
    vistos = set()

    for idioma in conhecidos:
        normal_idioma = _normalizar(idioma)
        if not re.search(rf"(?<!\w){re.escape(normal_idioma)}(?!\w)", normal_base):
            continue
        if normal_idioma in vistos:
            continue
        vistos.add(normal_idioma)

        nivel = None
        linha_idioma = next((linha for linha in linhas_secao if normal_idioma in _normalizar(linha)), "")
        contexto = _normalizar(linha_idioma)
        for chave, rotulo in _NIVEIS_IDIOMA.items():
            if chave in contexto:
                nivel = rotulo
                break
        resultado.append({"idioma": idioma, "nivel": nivel})
    return resultado[:20]


def analisar_curriculo(
    texto: str,
    habilidades_catalogo: Iterable[str] = (),
    idiomas_catalogo: Iterable[str] = (),
) -> dict:
    """Converte o texto extraído em uma prévia estruturada para revisão humana."""
    linhas = [_limpar_linha(linha) for linha in texto.splitlines() if _limpar_linha(linha)]
    secoes = _separar_secoes(linhas)
    cabecalho = secoes.get("cabecalho", [])
    cidade, estado = _localizacao(cabecalho + linhas[:15])

    titulo = next((linha for linha in cabecalho if _CARGO_RE.search(_normalizar(linha))), None)
    if not titulo:
        candidatos = [
            linha for linha in cabecalho[1:5]
            if "@" not in linha and "http" not in linha.lower() and not re.search(r"\d{8,}", linha)
        ]
        titulo = candidatos[0] if candidatos else None

    resumo_linhas = secoes.get("resumo", [])
    resumo = " ".join(resumo_linhas[:6]).strip()[:2000] or None

    experiencias = _parse_experiencias(secoes.get("experiencias", []))
    formacoes = _parse_formacoes(secoes.get("formacoes", []))
    certificados = _parse_formacoes(secoes.get("certificados", []), certificado=True)

    anos = [int(exp["data_inicio"][:4]) for exp in experiencias if exp.get("data_inicio")]
    experiencia_anos = None
    if anos:
        from datetime import date
        experiencia_anos = max(0, min(80, date.today().year - min(anos)))

    dados = {
        "versao": PARSER_VERSION,
        "contato": {
            "nome": cabecalho[0][:150] if cabecalho and "@" not in cabecalho[0] else None,
            "email": (re.search(r"[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}", texto).group(0) if re.search(r"[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}", texto) else None),
            "telefone": (re.search(r"(?:\+55\s*)?\(?\d{2}\)?[\s.-]*\d{4,5}[\s.-]*\d{4}", texto).group(0) if re.search(r"(?:\+55\s*)?\(?\d{2}\)?[\s.-]*\d{4,5}[\s.-]*\d{4}", texto) else None),
        },
        "perfil": {
            "titulo_profissional": titulo[:150] if titulo else None,
            "resumo": resumo,
            "cidade": cidade,
            "estado": estado,
            "linkedin_url": _url(texto, "linkedin.com/"),
            "github_url": _url(texto, "github.com/"),
            "portfolio_url": _portfolio(texto),
            "experiencia_anos": experiencia_anos,
        },
        "experiencias": experiencias,
        "formacoes": formacoes[:30],
        "certificados": certificados,
        "projetos": _parse_projetos(secoes.get("projetos", [])),
        "habilidades": _parse_habilidades(texto, secoes.get("habilidades", []), habilidades_catalogo),
        "idiomas": _parse_idiomas(texto, secoes.get("idiomas", []), idiomas_catalogo),
        "meta": {
            "linhas_analisadas": len(linhas),
            "secoes_identificadas": [secao for secao in secoes if secao != "cabecalho"],
            "avisos": [
                "Revise os dados antes de importar. Datas identificadas apenas pelo ano usam 1º de janeiro como referência."
            ],
        },
    }
    return dados
